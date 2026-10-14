import re
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace, MethodType
from unittest.mock import patch
import inspect
import pandas as pd
from openpyxl import load_workbook
from WallsEC2 import WallsEC2App, WallDesigner, clean_dataframe

class Variable:
    def __init__(self,value='0'): self.value=value
    def get(self): return self.value
    def set(self,value): self.value=value

class ExportTests(unittest.TestCase):
    def test_xlsx_composite_traceability(self):
        app=SimpleNamespace(update_idletasks=lambda:None)
        for name in re.findall(r'self\.(var_\w+)',inspect.getsource(WallsEC2App.export_excel)):
            setattr(app,name,Variable())
        app.progress_var=Variable(); app.status_var=Variable()
        values=dict(var_elu_cases='101',var_axial_convention='Compressão positiva',var_l0='3000',
                    var_phi_eff='1.5',var_eps_ud='.01',var_thickness='200',var_cover='30',
                    var_concrete='C30/37',var_steel='500')
        for name,value in values.items(): setattr(app,name,Variable(value))
        app.df_clean=clean_dataframe(pd.DataFrame([dict(panel=1,node=1,case=101,mxx=10,myy=10,mxy=0,
                                                       qxx=0,qyy=0,nxx=100,nyy=100,nxy=0)]))
        d=WallDesigner(200,30,'C30/37',500,elu_cases='101',axial_convention='Compressão positiva',
                       effective_length_mm=3000,creep_effective=1.5)
        app.df_results=d.design_dataframe(app.df_clean)
        for name in ['df_summary','df_governing','df_zones','df_optimization','df_diagnostic',
                     'df_data_validation','df_unit_check','df_detailed','df_notes']:
            setattr(app,name,pd.DataFrame())
        app._excel_sheet_description=MethodType(WallsEC2App._excel_sheet_description,app)
        with tempfile.TemporaryDirectory() as folder:
            target=str(Path(folder)/'results.xlsx')
            with patch('WallsEC2.filedialog.asksaveasfilename',return_value=target), \
                 patch('WallsEC2.messagebox.showinfo'),patch('WallsEC2.messagebox.showwarning'), \
                 patch('WallsEC2.messagebox.showerror',side_effect=AssertionError):
                WallsEC2App.export_excel(app)
            wb=load_workbook(target,data_only=True)
            self.assertIn('18_NM_Stability',wb.sheetnames)
            rows=list(wb['18_NM_Stability'].values)
            headers=rows[0]
            self.assertIn('NM_Y_util',headers)
            self.assertAlmostEqual(rows[1][headers.index('NYY_kN_m')],100)
            self.assertEqual(rows[1][headers.index('Axial_convention')],'Compressão positiva')
            self.assertEqual(rows[1][headers.index('Stability_l0_mm')],3000)
            wb.close()
